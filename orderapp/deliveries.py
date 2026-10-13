"""Morning delivery list: the email listing every order going out on the truck today.
Read it (from Outlook automatically, or pasted), match each line to an open order in the app by PO #,
sales order / invoice #, or customer name, and mark the matches Delivered."""
import io
import re
from datetime import datetime

from .db import OPEN_STATUSES, add_status, dumps, loads, norm_name, now, q, setting, x

DELIVERY_SCHEMA = """
CREATE TABLE IF NOT EXISTS delivery_checks (
    id INTEGER PRIMARY KEY, msg_id TEXT, source TEXT, subject TEXT, received TEXT, body TEXT,
    results TEXT, applied TEXT, created TEXT
);
"""
DONE = "Delivered"
_NUM = re.compile(r"[A-Za-z]{0,4}[-#: ]{0,2}\d[\w-]{2,}")


def _norm(s):
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def _digits(s):
    return re.sub(r"\D", "", s or "")


def attachment_text(attachments):
    """Text out of a PDF / CSV / Excel delivery sheet."""
    out = []
    for fn, data in attachments or []:
        low = fn.lower()
        try:
            if low.endswith(".pdf"):
                import fitz
                with fitz.open(stream=data, filetype="pdf") as d:
                    out.append("\n".join(p.get_text() for p in d))
            elif low.endswith((".csv", ".txt")):
                out.append(data.decode("utf-8-sig", errors="replace"))
            elif low.endswith((".xlsx", ".xlsm")):
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        out.append("  ".join(str(c) for c in row if c is not None))
        except Exception:  # noqa: BLE001 - a bad attachment shouldn't stop the rest
            continue
    return "\n".join(out)


def candidates():
    st = [s for s in OPEN_STATUSES if s != "Draft"]
    return q(f"""SELECT o.*, c.name cname FROM orders o LEFT JOIN customers c ON c.id=o.customer_id
                 WHERE o.status IN ({','.join('?' * len(st))}) ORDER BY o.id""", st)


def match(text):
    """[{order_id, cname, po, vendor_ref, status, how, sure, line}] - one entry per matched order."""
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    tokens = {}
    for ln in lines:
        for t in _NUM.findall(ln):
            n = _norm(t)
            tokens.setdefault(n, ln)
            d = _digits(t)
            if len(d) >= 4:
                tokens.setdefault("#" + d, ln)
    found = {}
    orders = candidates()
    for o in orders:
        for field, how in (("po", "PO #"), ("vendor_ref", "Sales order #")):
            v = o[field] or ""
            if not v.strip():
                continue
            n, d = _norm(v), _digits(v)
            line = tokens.get(n) or (tokens.get("#" + d) if len(d) >= 4 else None)
            if line:
                found[o["id"]] = {"order_id": o["id"], "cname": o["cname"] or "", "po": o["po"] or "",
                                  "vendor_ref": o["vendor_ref"] or "", "status": o["status"], "title": o["title"],
                                  "how": how, "sure": True, "line": line[:140]}
                break
    # a line whose PO / SO # already matched is used up; other lines naming a customer = "maybe" for each of
    # that customer's open orders (you confirm - names alone aren't proof)
    used = {v["line"] for v in found.values()}
    by_cust = {}
    for o in orders:
        if o["id"] in found or not o["customer_id"] or not o["cname"]:
            continue
        by_cust.setdefault(o["customer_id"], []).append(o)
    for cid, os_ in by_cust.items():
        key = norm_name(os_[0]["cname"])
        if len(key) < 5:
            continue
        hit = next((ln for ln in lines if ln[:140] not in used and key in norm_name(ln)), None)
        if not hit:
            continue
        for o in os_:
            found[o["id"]] = {"order_id": o["id"], "cname": o["cname"], "po": o["po"] or "",
                              "vendor_ref": o["vendor_ref"] or "", "status": o["status"], "title": o["title"],
                              "how": "Customer name", "sure": False, "line": hit[:140]}
    return sorted(found.values(), key=lambda r: (not r["sure"], r["cname"]))


def unmatched_lines(text, results):
    used = {r["line"] for r in results}
    out = []
    for ln in (text or "").splitlines():
        s = ln.strip()
        if s and s[:140] not in used and (_NUM.search(s) or len(s) > 12) and not s.lower().startswith(
                ("from:", "sent:", "to:", "subject:", "cc:")):
            out.append(s[:140])
    return out[:60]


def apply(order_ids, when=None):
    when = when or datetime.now().strftime("%m/%d/%Y")
    n = 0
    for oid in order_ids:
        o = q("SELECT status FROM orders WHERE id=?", (oid,), one=True)
        if not o or o["status"] == DONE:
            continue
        x("UPDATE orders SET status=?, updated=? WHERE id=?", (DONE, now(), oid))
        add_status(oid, DONE, f"On the delivery list {when}")
        n += 1
    return n


def record(text, source, subject="", received="", msg_id=None, auto_apply=False):
    """Match a delivery list, save the check, and (auto) apply the sure matches. Returns (check id, results, applied)."""
    res = match(text)
    applied = apply([r["order_id"] for r in res if r["sure"]]) if auto_apply else 0
    cid = x("""INSERT INTO delivery_checks(msg_id, source, subject, received, body, results, applied, created)
               VALUES (?,?,?,?,?,?,?,?)""",
            (msg_id, source, subject, received, text[:200000], dumps(res),
             dumps([r["order_id"] for r in res if r["sure"]] if auto_apply else []), now()))
    return cid, res, applied


def check_email(prov=None):
    """Find today's delivery-list email(s) in Outlook and process new ones. Never sends anything."""
    from .emailer import provider
    prov = prov or provider(setting)
    word = (setting("delivery_subject_word") or "Deliver").strip()
    out = {"found": 0, "new": 0, "matched": 0, "applied": 0, "error": ""}
    if prov.name == "eml":
        out["error"] = "Email isn't connected (Settings -> Email)."
        return out
    ok, why = prov.available()
    if not ok:
        out["error"] = why
        return out
    try:
        msgs = prov.list_messages(days=1, limit=20, subject_word=word)
    except Exception as e:  # noqa: BLE001
        out["error"] = f"Couldn't read the inbox: {e}"
        return out
    out["found"] = len(msgs)
    seen = {r["msg_id"] for r in q("SELECT msg_id FROM delivery_checks WHERE msg_id IS NOT NULL")}
    auto = (setting("delivery_auto") or "1") == "1"
    from .order_notices import is_notice
    for m in msgs:
        if m["msg_id"] in seen:
            continue
        if is_notice(m.get("subject")):
            continue                    # an order-received email, not the truck list
        try:
            full = prov.get_message(m["msg_id"])
        except Exception:  # noqa: BLE001
            continue
        text = (full.get("body") or "") + "\n" + attachment_text(full.get("attachments"))
        _, res, applied = record(text, "email", full.get("subject", ""), full.get("received", ""), m["msg_id"], auto)
        out["new"] += 1
        out["matched"] += len(res)
        out["applied"] += applied
    return out


def recent(limit=10, kind="delivery"):
    rows = []
    for r in q("SELECT * FROM delivery_checks WHERE coalesce(kind,'delivery')=? ORDER BY id DESC LIMIT ?",
               (kind, limit)):
        d = dict(r)
        d["results"] = loads(r["results"], [])
        d["applied"] = loads(r["applied"], [])
        rows.append(d)
    return rows


def pending_maybes():
    """Name-only matches from recent lists that are still open (for the dashboard badge)."""
    ids = set()
    for c in recent(5):
        for r in c["results"]:
            if not r["sure"]:
                o = q("SELECT status FROM orders WHERE id=?", (r["order_id"],), one=True)
                if o and o["status"] != DONE:
                    ids.add(r["order_id"])
    return len(ids)
