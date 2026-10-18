"""Export the app's data as Excel or CSV (Customers, contacts, units & sizes, orders, sales, CRM, items).

The customer export uses the same column names the Import CSV page understands, so a list can be
exported, filled in / fixed in Excel and imported back (Review shows every change before it is saved).
"""
import csv
import io
import json
from datetime import date

from flask import Blueprint, Response, abort, render_template, request

from ..db import q
from ..security import login_required

bp = Blueprint("exports", __name__)


def _customers(args):
    sql = """SELECT c.name, coalesce(nullif(c.acct_type,''),'Property'), c.services, c.acct, c.address, c.city, c.state, c.zip, c.mgmt, c.phone, c.email, c.status, c.tags,
                    (SELECT COUNT(*) FROM contacts k WHERE k.customer_id=c.id) AS n_contacts,
                    (SELECT COUNT(*) FROM units u WHERE u.customer_id=c.id) AS n_units,
                    (SELECT COUNT(*) FROM orders o WHERE o.customer_id=c.id) AS n_orders,
                    (SELECT MAX(o.order_date) FROM orders o WHERE o.customer_id=c.id) AS last_order,
                    (SELECT ROUND(SUM(s.total), 2) FROM sales s WHERE s.customer_id=c.id) AS sales_total,
                    (SELECT MAX(s.invoice_date) FROM sales s WHERE s.customer_id=c.id) AS last_invoice,
                    (SELECT MAX(a.occurred_at) FROM activities a WHERE a.customer_id=c.id) AS last_contact,
                    c.notes
             FROM customers c"""
    where = []
    if args.get("missing") == "address":
        where.append("coalesce(trim(c.address),'')=''")
    elif args.get("missing") == "mgmt":
        where.append("coalesce(trim(c.mgmt),'')=''")
    elif args.get("missing") == "any":
        where.append("(coalesce(trim(c.address),'')='' OR coalesce(trim(c.mgmt),'')='')")
    if args.get("status"):
        where.append("coalesce(c.status,'')=?")
    from ..accounts import ACCOUNT_TYPES
    typ = args.get("type") if args.get("type") in ACCOUNT_TYPES else ""
    if typ:
        where.append("coalesce(nullif(c.acct_type,''),'Property')=?")
    sql += (" WHERE " + " AND ".join(where)) if where else ""
    sql += " ORDER BY c.name COLLATE NOCASE"
    params = ((args["status"],) if args.get("status") else ()) + ((typ,) if typ else ())
    rows = q(sql, params)
    head = ["Property", "Account Type", "Services", "Acct #", "Address", "City", "State", "Zip", "Mgmt Co", "Phone", "Email", "Status", "Tags",
            "Contacts", "Units", "Orders", "Last Order", "Sales Total", "Last Invoice", "Last Contact", "Notes"]
    return head, [list(r) for r in rows]


def _contacts(args):
    rows = q("""SELECT c.name, c.acct, k.name, k.role, k.phone, k.email, k.notes, c.mgmt, c.city
                FROM contacts k JOIN customers c ON c.id=k.customer_id
                ORDER BY c.name COLLATE NOCASE, k.name COLLATE NOCASE""")
    return (["Property", "Acct #", "Contact Name", "Contact Role", "Contact Phone", "Contact Email", "Notes",
             "Mgmt Co", "City"], [list(r) for r in rows])


SIZE_FIELDS = ["width", "height", "qty", "color", "style", "swing", "core", "item_no", "description"]


def _sizes(args):
    rows = q("""SELECT c.name, c.acct, u.unit_number, u.building, f.name AS fp, m.room, m.product, m.data,
                       m.verified, m.notes, m.updated
                FROM measurements m JOIN customers c ON c.id=m.customer_id
                LEFT JOIN units u ON u.id=m.unit_id LEFT JOIN floorplans f ON f.id=m.floorplan_id
                ORDER BY c.name COLLATE NOCASE, u.unit_number, m.product""")
    out = []
    for r in rows:
        try:
            d = json.loads(r["data"] or "{}")
        except ValueError:
            d = {}
        out.append([r["name"], r["acct"], r["unit_number"], r["building"], r["fp"], r["room"], r["product"]] +
                   [d.get(k, "") for k in SIZE_FIELDS] + ["Yes" if r["verified"] else "", r["notes"], r["updated"]])
    return (["Property", "Acct #", "Unit", "Building", "Floorplan", "Room", "Product", "Width", "Height", "Qty",
             "Color", "Style", "Swing", "Core", "Item #", "Description", "Verified", "Notes", "Updated"], out)


def _orders(args):
    from ..catalog import FORMS
    from ..order_kind import status_label
    rows = q("""SELECT o.id, o.order_date, c.name, c.acct, o.form_key, o.title, o.kind, o.po, o.status,
                       o.est_value, o.due_date, o.sent_at, o.sent_to, o.source, o.notes
                FROM orders o LEFT JOIN customers c ON c.id=o.customer_id ORDER BY o.id DESC""")
    out = [[r["id"], r["order_date"], r["name"], r["acct"],
            FORMS.get(r["form_key"], {}).get("title", r["form_key"] or ""), r["title"],
            "Estimate" if (r["kind"] or "") == "quote" else "Order", r["po"],
            status_label(r["status"], r["kind"] or "order"), r["est_value"], r["due_date"], r["sent_at"],
            r["sent_to"], r["source"], r["notes"]] for r in rows]
    return (["Order #", "Date", "Property", "Acct #", "Form", "Title", "Order / Estimate", "PO", "Status",
             "Est. Value", "Due", "Sent", "Sent To", "Source", "Notes"], out)


def _sales(args):
    rows = q("""SELECT s.invoice_date, s.invoice_no, coalesce(c.name, s.customer_name), s.acct, s.po, s.rep, s.total,
                       s.cost FROM sales s LEFT JOIN customers c ON c.id=s.customer_id
                ORDER BY s.invoice_date DESC, s.invoice_no""")
    return (["Invoice Date", "Invoice #", "Property", "Acct #", "PO", "Rep", "Total", "Cost"], [list(r) for r in rows])


def _activities(args):
    rows = q("""SELECT a.occurred_at, c.name, a.contact_name, k.role, a.kind, a.subject, a.outcome, a.location,
                       a.topics, a.notes FROM activities a LEFT JOIN customers c ON c.id=a.customer_id
                LEFT JOIN contacts k ON k.id=a.contact_id ORDER BY a.occurred_at DESC""")
    return (["When", "Property", "Talked To", "Position", "How", "Subject", "Warm / Cold", "Location", "Topics",
             "Notes"], [list(r) for r in rows])


def _tasks(args):
    rows = q("""SELECT t.due_date, t.title, c.name, CASE WHEN t.done THEN 'Done' ELSE 'Open' END, t.done_at, t.created
                FROM tasks t LEFT JOIN customers c ON c.id=t.customer_id ORDER BY t.done, t.due_date""")
    return (["Due", "Task", "Property", "Status", "Done On", "Created"], [list(r) for r in rows])


def _deals(args):
    rows = q("""SELECT d.name, c.name, d.product, d.value, d.stage, d.close_date, d.notes, d.updated
                FROM deals d LEFT JOIN customers c ON c.id=d.customer_id ORDER BY d.stage, d.close_date""")
    return (["Deal", "Property", "Product", "Value", "Stage", "Close Date", "Notes", "Updated"],
            [list(r) for r in rows])


def _items(args):
    rows = q("""SELECT item_no, description, vendor, category, list_price, avg_price, avg_cost, margin_pct, on_hand,
                       on_so, available, on_po, reorder_pt, next_deliv, sales_week, qty_sold, last_sold
                FROM items ORDER BY item_no""")
    return (["Item #", "Description", "Vendor", "Category", "List Price", "Avg Price", "Avg Cost", "Margin %",
             "On Hand", "On Sales Order", "Available", "On PO", "Reorder Point", "Next Delivery", "Sales / Week",
             "Qty Sold", "Last Sold"], [list(r) for r in rows])


def _holds(args):
    rows = q("SELECT * FROM account_holds")
    if not rows:
        return ["Account"], []
    keys = rows[0].keys()
    return [k.replace("_", " ").title() for k in keys], [list(r) for r in rows]


DATASETS = {
    "customers": ("Customer list", "Every property / account with address, mgmt co., counts, sales total and "
                  "last contact. Same columns the Import CSV page reads, so you can fix it in Excel and import it "
                  "back.", _customers),
    "contacts": ("Contacts (people & positions)", "Who works at each property - name, position, phone, email.",
                 _contacts),
    "sizes": ("Units & sizes", "Every saved size by property and unit.", _sizes),
    "orders": ("Orders & estimates", "Every order / estimate with its form, status, PO and dates.", _orders),
    "sales": ("Sales invoices", "Imported invoices (My Sales).", _sales),
    "activities": ("CRM contact log", "Every logged visit, call, text and email.", _activities),
    "tasks": ("CRM tasks", "Follow-ups and to-dos.", _tasks),
    "deals": ("Pipeline deals", "Open and closed deals.", _deals),
    "items": ("Items & stock", "Item prices, stock, margins.", _items),
    "holds": ("Account holds", "Accounts on hold.", _holds),
}


def _cell(v):
    return "" if v is None else v


def to_csv(head, rows):
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(head)
    for r in rows:
        w.writerow([_cell(v) for v in r])
    return ("﻿" + out.getvalue()).encode("utf-8")


def to_xlsx(sheets):
    """sheets: [(title, head, rows)] -> xlsx bytes (bold frozen header, filters, sensible widths)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for title, head, rows in sheets:
        ws = wb.create_sheet(title[:31].replace("/", "-").replace("&", "and"))
        ws.append(head)
        for r in rows:
            ws.append([_cell(v) for v in r])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A68")
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(rows) + 1}"
        for i, h in enumerate(head, 1):
            longest = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:300]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@bp.route("/export")
@login_required
def index():
    from ..db import CUSTOMER_STATUSES
    counts = {
        "customers": q("SELECT COUNT(*) n FROM customers", one=True)["n"],
        "no_address": q("SELECT COUNT(*) n FROM customers WHERE coalesce(trim(address),'')=''", one=True)["n"],
        "no_mgmt": q("SELECT COUNT(*) n FROM customers WHERE coalesce(trim(mgmt),'')=''", one=True)["n"],
    }
    from ..accounts import ACCOUNT_TYPES
    return render_template("export.html", datasets=DATASETS, counts=counts, statuses=CUSTOMER_STATUSES,
                           acct_types=ACCOUNT_TYPES)


@bp.route("/export/<name>.<fmt>")
@login_required
def download(name, fmt):
    stamp = date.today().isoformat()
    if name == "everything" and fmt == "xlsx":
        sheets = [(DATASETS[k][0], *DATASETS[k][2]({})) for k in DATASETS]
        return Response(to_xlsx(sheets), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f"attachment; filename=\"AIS Sales Support - all data {stamp}.xlsx\""})
    if name not in DATASETS or fmt not in ("csv", "xlsx"):
        abort(404)
    label, _desc, fn = DATASETS[name]
    head, rows = fn(request.args)
    extra = {"address": " - missing address", "mgmt": " - missing mgmt co", "any": " - missing info"}.get(
        request.args.get("missing", ""), "")
    fname = f"{label}{extra} {stamp}"
    if fmt == "csv":
        return Response(to_csv(head, rows), mimetype="text/csv",
                        headers={"Content-Disposition": f"attachment; filename=\"{fname}.csv\""})
    return Response(to_xlsx([(label, head, rows)]),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=\"{fname}.xlsx\""})
