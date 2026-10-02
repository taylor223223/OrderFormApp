"""CSV / Excel import -> proposed changes on the Review page."""
import csv
import io
import re

from .catalog import FORMS, PRODUCTS, forms_for_product
from .changes import add_change, new_batch, propose_contact, propose_customer_fields
from .db import dumps, find_customer, loads, norm_name, q

# target -> header synonyms (normalised: lowercase alnum only)
SYNONYMS = {
    "name": ["name", "customer", "customername", "property", "propertyname", "community", "account name",
             "accountname", "site", "complex"],
    "acct": ["acct", "acct#", "account", "accountnumber", "accountno", "acctno", "acctnumber", "custno",
             "customernumber", "customerid"],
    "address": ["address", "addr", "street", "streetaddress", "address1"],
    "city": ["city", "town"],
    "state": ["state", "st", "province"],
    "zip": ["zip", "zipcode", "postal", "postalcode"],
    "mgmt": ["mgmt", "mgmtco", "management", "managementcompany", "mgmtcompany", "pmc", "owner"],
    "phone": ["phone", "phonenumber", "officephone", "tel", "telephone", "propertyphone"],
    "email": ["email", "emailaddress", "propertyemail", "officeemail"],
    "contact_name": ["contact", "contactname", "manager", "propertymanager", "maintenancesupervisor",
                     "supervisor", "orderedby", "requestedby"],
    "contact_email": ["contactemail", "manageremail", "supervisoremail"],
    "contact_phone": ["contactphone", "cell", "mobile", "managerphone", "supervisorphone"],
    "contact_role": ["contactrole", "role", "title", "position"],
    "unit": ["unit", "unit#", "unitnumber", "unitno", "apt", "apartment", "aptno"],
    "building": ["building", "bldg", "buildingno"],
    "floorplan": ["floorplan", "floorplanname", "plan", "layout", "unittype", "fp", "model"],
    "product": ["product", "producttype", "item", "itemtype", "category", "type", "description", "desc"],
    "room": ["room", "location", "loc", "area"],
    "qty": ["qty", "quantity", "count", "pcs"],
    "item_no": ["itemno", "item#", "itemnumber", "sku", "partno", "part#", "partnumber"],
    "width": ["width", "w", "doorwidth", "windowwidth", "actualwidth"],
    "height": ["height", "h", "doorheight", "windowheight", "actualheight"],
    "length": ["length", "len", "lf", "linearft", "linearfeet"],
    "depth": ["depth", "d"],
    "color": ["color", "colour"],
    "finish": ["finish", "material"],
    "style": ["style", "profile", "doorstyle", "panel"],
    "core": ["core", "coretype"],
    "swing": ["swing", "hand", "handing", "doorswing"],
    "thickness": ["thickness", "doorthickness"],
    "po": ["po", "po#", "ponumber", "purchaseorder", "pono"],
    "order_date": ["date", "orderdate", "dateordered", "orderedon"],
    "status": ["status", "orderstatus"],
    "notes": ["notes", "note", "comments", "comment", "memo"],
    "verified": ["verified", "preverified", "confirmed"],
}
CUSTOMER_TARGETS = ["name", "acct", "address", "city", "state", "zip", "mgmt", "phone", "email"]
CONTACT_TARGETS = ["contact_name", "contact_email", "contact_phone", "contact_role"]
META_TARGETS = ["unit", "building", "floorplan", "product", "room", "po", "order_date", "status", "notes",
                "verified"]


def all_targets():
    """(value, label) list for the column-mapping dropdowns."""
    t = [("", "(ignore)")]
    t += [(k, "Customer: " + k) for k in CUSTOMER_TARGETS]
    t += [(k, "Contact: " + k.split("_", 1)[1]) for k in CONTACT_TARGETS]
    t += [(k, "Unit/Order: " + k) for k in META_TARGETS]
    seen = set(CUSTOMER_TARGETS + CONTACT_TARGETS + META_TARGETS)
    for fk in FORMS:
        for f in FORMS[fk]["blocks"][0]:
            if f["key"] not in seen and f["key"] not in ("unit", "line_comments", "other_mods"):
                seen.add(f["key"])
                t.append((f["key"], "Measurement: " + f["label"]))
    for extra in ("length", "depth", "color", "finish", "hardware"):
        if extra not in seen:
            seen.add(extra)
            t.append((extra, "Measurement: " + extra))
    return t


def _n(s):
    return re.sub(r"[^a-z0-9#]", "", str(s or "").lower())


def guess_mapping(headers):
    m = {}
    used = set()
    keys = {k for k, _ in all_targets() if k}
    for h in headers:
        hn = _n(h)
        target = ""
        for k, syns in SYNONYMS.items():
            if hn in [_n(s) for s in syns]:
                target = k
                break
        if not target:
            # match a measurement key or its label: "Door Location", "door_location", "Slat Width"...
            for k, label in all_targets():
                if not k:
                    continue
                lab = label.split(": ", 1)[-1]
                if hn.replace("#", "") in (_n(k), _n(lab), _n(lab.split(" (")[0])):
                    target = k
                    break
        if target in used and target not in ("notes",):
            target = ""
        if target:
            used.add(target)
        m[h] = target
    return m


def read_table(raw: bytes, filename: str):
    """Return (headers, rows[list of dict])."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = []
        for row in it:
            if row and any(c not in (None, "") for c in row):
                headers = [str(c).strip() if c is not None else f"Column{i + 1}" for i, c in enumerate(row)]
                break
        rows = []
        for row in it:
            if row and any(c not in (None, "") for c in row):
                rows.append({headers[i]: ("" if c is None else str(c).strip()) for i, c in enumerate(row)
                             if i < len(headers)})
        return headers, rows
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rdr = csv.reader(io.StringIO(text), dialect)
    rows_raw = [r for r in rdr if any(c.strip() for c in r)]
    if not rows_raw:
        return [], []
    headers = [h.strip() or f"Column{i + 1}" for i, h in enumerate(rows_raw[0])]
    rows = [{headers[i]: c.strip() for i, c in enumerate(r) if i < len(headers)} for r in rows_raw[1:]]
    return headers, rows


PRODUCT_WORDS = [
    ("prehung_door", ["pre-hung", "prehung", "pre hung"]),
    ("bypass_door", ["bypass", "by-pass", "bi-pass", "bipass", "sliding closet"]),
    ("screen_door", ["screen door", "security door", "patio screen", "sliding screen"]),
    ("window_screen", ["window screen", "screen"]),
    ("vertical_blind", ["vertical", "vert blind", "verts"]),
    ("horizontal_blind", ["horizontal", "mini blind", "miniblind", "faux wood", "blind"]),
    ("garage_door", ["garage"]),
    ("entry_door", ["entry", "front door", "exterior door", "ext door", "steel door"]),
    ("storage_door", ["storage", "utility", "water heater", "mechanical"]),
    ("closet_door", ["closet", "bifold", "bi-fold"]),
    ("cabinet_door", ["cabinet door", "drawer front", "cab door"]),
    ("cabinet", ["cabinet", "vanity"]),
    ("baseboard", ["baseboard", "base board", "base trim", "shoe mold"]),
    ("interior_door", ["interior door", "int door", "bedroom door", "bath door", "door"]),
]


def guess_product(text):
    t = str(text or "").lower()
    if not t:
        return None
    for key, label in ((k, v["label"].lower()) for k, v in PRODUCTS.items()):
        if t == key or t == label:
            return key
    for key, words in PRODUCT_WORDS:
        if any(w in t for w in words):
            return key
    return None


def _meas_exists(cid, fp_name, unit_no, product, room, data):
    if not cid:
        return False
    rows = q("""SELECT m.* FROM measurements m LEFT JOIN floorplans f ON f.id=m.floorplan_id
                LEFT JOIN units u ON u.id=m.unit_id
                WHERE m.customer_id=? AND m.product=? AND lower(coalesce(m.room,''))=lower(?)
                AND (lower(coalesce(f.name,''))=lower(?) OR lower(coalesce(u.unit_number,''))=lower(?))""",
             (cid, product, room or "", fp_name or "", unit_no or ""))
    for r in rows:
        if loads(r["data"]) == data:
            return True
    return False


def build_batch(rows, mapping, filename, default_customer_id=None):
    """Turn mapped rows into a review batch. Returns (batch_id, summary)."""
    batch = new_batch("csv", f"Import: {filename}")
    meas_keys = {k for k, _ in all_targets() if k} - set(CUSTOMER_TARGETS) - set(CONTACT_TARGETS) - \
        set(META_TARGETS)
    groups = {}   # key -> dict
    order = []
    default_cust = q("SELECT * FROM customers WHERE id=?", (default_customer_id,), one=True) \
        if default_customer_id else None

    for row in rows:
        v = {}
        for col, target in mapping.items():
            if target and row.get(col, "") != "":
                v[target] = (v[target] + " " + row[col]) if target == "notes" and v.get(target) else row[col]
        cust_fields = {k: v[k] for k in CUSTOMER_TARGETS if v.get(k)}
        cust = find_customer(cust_fields.get("name"), cust_fields.get("acct"), cust_fields.get("email"))
        if cust is None and not cust_fields.get("name") and default_cust is not None:
            cust = default_cust
        if cust is not None:
            gkey = f"id:{cust['id']}"
        elif cust_fields.get("name") or cust_fields.get("acct"):
            gkey = "new:" + (norm_name(cust_fields.get("name")) or ("acct" + cust_fields.get("acct", "")))
        else:
            continue
        g = groups.get(gkey)
        if g is None:
            g = {"cust": cust, "fields": {}, "contacts": [], "units": {}, "fps": set(), "meas": [], "orders": {}}
            groups[gkey] = g
            order.append(gkey)
        for k, val in cust_fields.items():
            g["fields"].setdefault(k, val)
        # notes on a row with no product/sizes are about the customer (delivery notes etc.)
        row_has_meas = any(v.get(k) for k in meas_keys) and (v.get("product") or v.get("item_no"))
        if v.get("notes") and not row_has_meas and v["notes"] not in g["fields"].get("notes", ""):
            g["fields"]["notes"] = "; ".join(x for x in [g["fields"].get("notes"), v["notes"]] if x)
        contact = {"name": v.get("contact_name", ""), "email": v.get("contact_email", ""),
                   "phone": v.get("contact_phone", ""), "role": v.get("contact_role", "")}
        if any(contact.values()) and contact not in g["contacts"]:
            g["contacts"].append(contact)
        fp = v.get("floorplan", "").strip()
        unit = v.get("unit", "").strip()
        if fp:
            g["fps"].add(fp)
        if unit:
            g["units"].setdefault(unit, {"floorplan": fp, "building": v.get("building", "")})
            if fp and not g["units"][unit]["floorplan"]:
                g["units"][unit]["floorplan"] = fp
        data = {k: v[k] for k in meas_keys if v.get(k)}
        product = guess_product(v.get("product")) or (guess_product(v.get("item_no")) if data else None)
        if data and (product or v.get("product")):
            product = product or "other"
            if product == "other" and v.get("product"):
                data.setdefault("style", v["product"])
            m = {"product": product, "room": v.get("room", ""), "data": data, "floorplan_name": fp,
                 "unit_number": unit, "building": v.get("building", ""), "notes": v.get("notes", ""),
                 "verified": str(v.get("verified", "")).lower() in ("y", "yes", "true", "1", "x"),
                 "_from_order": bool(v.get("po") or v.get("order_date"))}
            # one saved measurement per layout/unit + product + room: the first row wins
            # (later rows for the same spot - e.g. past orders - still become order history)
            slot = (fp.lower() if fp else "unit:" + unit.lower(), product, (v.get("room") or "").lower())
            if slot not in g.setdefault("slots", set()):
                g["slots"].add(slot)
                g["meas"].append(m)
            if v.get("po") or v.get("order_date"):
                okey = (v.get("po", ""), v.get("order_date", ""))
                o = g["orders"].setdefault(okey, {"po": v.get("po", ""), "order_date": v.get("order_date", ""),
                                                  "status": v.get("status", "") or "Completed",
                                                  "products": [], "blocks": []})
                o["products"].append(product)
                blk = dict(data)
                if unit:
                    blk["unit"] = unit
                if v.get("room"):
                    blk["room"] = v["room"]
                o["blocks"].append(blk)

    summary = {"customers_new": 0, "fields": 0, "contacts": 0, "units": 0, "measurements": 0, "orders": 0}
    for gkey in order:
        g = groups[gkey]
        cust = g["cust"]
        ref = None
        if cust is None:
            ref = gkey
            add_change(batch, "new_customer", customer_ref=ref, payload=g["fields"],
                       new=g["fields"].get("name") or g["fields"].get("acct"))
            summary["customers_new"] += 1
        else:
            summary["fields"] += propose_customer_fields(batch, cust, g["fields"])
        cid = cust["id"] if cust is not None else None
        for c in g["contacts"]:
            summary["contacts"] += propose_contact(batch, cust, c, ref)
        for fp in sorted(g["fps"]):
            if cid and q("SELECT 1 FROM floorplans WHERE customer_id=? AND lower(name)=lower(?)", (cid, fp), one=True):
                continue
            add_change(batch, "new_floorplan", customer_id=cid, customer_ref=ref, payload={"name": fp}, new=fp)
        for un, info in g["units"].items():
            ex = cid and q("SELECT * FROM units WHERE customer_id=? AND lower(unit_number)=lower(?)",
                           (cid, un), one=True)
            if ex and (ex["floorplan_id"] or not info["floorplan"]):
                continue
            add_change(batch, "new_unit", customer_id=cid, customer_ref=ref,
                       payload={"unit_number": un, "building": info["building"], "floorplan_name": info["floorplan"]},
                       new=un)
            summary["units"] += 1
        for m in g["meas"]:
            from_order = m.pop("_from_order", False)
            if from_order and m["unit_number"] and not m["floorplan_name"]:
                # a past-order line for a unit whose floorplan already has this product measured:
                # keep it as order history only, don't save a duplicate unit measurement
                ufp = g["units"].get(m["unit_number"], {}).get("floorplan", "")
                if not ufp and cid:
                    rowu = q("""SELECT f.name FROM units u JOIN floorplans f ON f.id=u.floorplan_id
                                WHERE u.customer_id=? AND lower(u.unit_number)=lower(?)""", (cid, m["unit_number"]), one=True)
                    ufp = rowu["name"] if rowu else ""
                if ufp and (any(o is not m and o["floorplan_name"].lower() == ufp.lower() and o["product"] == m["product"]
                                for o in g["meas"]) or (cid and q(
                        """SELECT 1 FROM measurements m JOIN floorplans f ON f.id=m.floorplan_id
                           WHERE m.customer_id=? AND lower(f.name)=lower(?) AND m.product=?""",
                        (cid, ufp, m["product"]), one=True))):
                    continue
            if _meas_exists(cid, m["floorplan_name"], m["unit_number"], m["product"], m["room"], m["data"]):
                continue
            add_change(batch, "new_measurement", customer_id=cid, customer_ref=ref, payload=m)
            summary["measurements"] += 1
        for (po, odate), o in g["orders"].items():
            if cid and po and q("SELECT 1 FROM orders WHERE customer_id=? AND po=?", (cid, po), one=True):
                continue
            prod = o["products"][0] if o["products"] else None
            fk = (forms_for_product(prod) or [None])[0] if prod else None
            header = {k: g["fields"].get(k, "") or ((cust[k] or "") if cust is not None else "")
                      for k in CUSTOMER_TARGETS if k != "email"}
            header.update({"po": po, "date": odate})
            title = FORMS[fk]["title"] if fk else "Imported order"
            add_change(batch, "history_order", customer_id=cid, customer_ref=ref,
                       payload={"form_key": fk, "po": po, "order_date": odate, "status": o["status"],
                                "title": title, "data": {"header": header, "blocks": o["blocks"]},
                                "source": "csv"},
                       new=f"PO {po}" if po else odate)
            summary["orders"] += 1
    return batch, summary


def save_upload_meta(path, filename, headers, mapping):
    return dumps({"path": path, "filename": filename, "headers": headers, "mapping": mapping})
