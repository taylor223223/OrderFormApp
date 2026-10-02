"""Purchase history: invoice-style Excel import -> history tab, search, keep-on-file."""
import io
import os
import re
import tempfile

import openpyxl
import pytest

_tmp = tempfile.mkdtemp()
os.environ["ORDERAPP_DATA"] = _tmp

from orderapp import create_app  # noqa: E402


@pytest.fixture(scope="module")
def client():
    app = create_app(db_path=os.path.join(_tmp, "hist.db"), testing=True)
    c = app.test_client()
    c.app = app
    return c


def csrf(c, path="/"):
    return re.search(r'name="csrf-token" content="([^"]+)"', c.get(path).get_data(as_text=True)).group(1)


def post(c, path, data=None, page="/", **kw):
    data = dict(data or {})
    data["csrf_token"] = csrf(c, page)
    return c.post(path, data=data, **kw)


def test_parse_description():
    from orderapp.importer import parse_description
    assert parse_description("VERT BLIND 96 x 84 ALABASTER") == {"width": "96", "height": "84", "color": "Alabaster"}
    assert parse_description('Window screen 35-1/2" X 47 3/4" bronze')["width"] == "35-1/2"


def _xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws["A1"] = "nothing here"
    d = wb.create_sheet("Sales Detail")
    d.append(["Apartment Interior Supply"])
    d.append(["Sales by Item Detail - Jan-Sep 2026"])
    d.append([])
    d.append(["Customer", "Unit", "Item", "Description", "Qty", "Invoice #", "Invoice Date", "Rate", "Amount"])
    import datetime
    d.append(["Tomscot Apartments", "104", "D-3080", "Int door 30x80 HC colonist primed", 2, "INV-1001",
              datetime.datetime(2026, 3, 14), 64.5, 129.0])
    d.append(["Tomscot Apartments", "206", "VB-9684", "Vertical blind 96 x 84 alabaster", 1, "INV-1001",
              datetime.datetime(2026, 3, 14), 118.0, 118.0])
    d.append(["Tomscot Apartments", "", "WS-BRZ", "Window screen 35-1/2 x 47-3/4 bronze", 3, "INV-1050",
              datetime.datetime(2026, 5, 2), 22.0, 66.0])
    d.append(["Highland", "12", "BP-48", "Bypass door 48 x 80 white", 1, "INV-2001",
              datetime.datetime(2026, 6, 9), 140.0, 140.0])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_history_import_and_search(client):
    c = client
    post(c, "/setup", {"username": "taylor", "password": "secret123", "password2": "secret123"}, page="/setup")
    r = post(c, "/customers/new", {"name": "Tomscot"})          # existing, named without "Apartments"
    tid = int(r.headers["Location"].rsplit("/", 1)[1])
    assert c.get("/import/template.csv").data.startswith("﻿Property".encode())

    r = post(c, "/import", {"file": (io.BytesIO(_xlsx()), "sales detail.xlsx")}, page="/import",
             content_type="multipart/form-data")
    html = r.get_data(as_text=True)
    token = re.search(r'name="token" value="([^"]+)"', html).group(1)
    sel = {}
    for i in range(9):
        m = re.search(rf'name="map_{i}".*?<option value="([^"]*)" selected', html, re.S)
        sel[f"map_{i}"] = m.group(1) if m else ""
    # header row found under the report title; columns recognised
    assert [sel[f"map_{i}"] for i in range(9)] == ["name", "unit", "product", "description", "qty", "invoice",
                                                   "order_date", "price", "amount"]
    r = post(c, "/import/apply_mapping", {"token": token, "filename": "sales detail.xlsx", **sel}, page="/import")
    bid = int(r.headers["Location"].rsplit("/", 1)[1])
    page = c.get(f"/review/{bid}").get_data(as_text=True)
    assert "NEW: Highland" in page
    ids = re.findall(r'name="accept" value="(\d+)"', page)
    post(c, f"/review/{bid}", {"accept": ids}, page=f"/review/{bid}")

    h = c.get(f"/customers/{tid}?tab=history").get_data(as_text=True)
    assert "Purchase History" in h and "INV-1001" in h and "2026-03-14" in h
    assert "96 x 84" in h and "Alabaster" in h and "Vertical Blind" in h and "35-1/2 x 47-3/4" in h
    assert "Window Screen" in h and "Bronze" in h
    # unit rows became on-file sizes for that unit
    j = c.get(f"/api/unit?customer={tid}&form=vertical_blind&unit=206").get_json()
    assert j["blocks"] and j["blocks"][0]["width"] == "96"
    # global search
    s = c.get("/history?q=bypass white").get_data(as_text=True)
    assert "Highland" in s and "48 x 80" in s and "Tomscot" not in s.split("<table")[-1]
    # order view lists lines for imported orders
    with c.app.app_context():
        from orderapp.db import q
        o = q("SELECT id FROM orders WHERE po='INV-1050'", one=True)
    ov = c.get(f"/orders/{o['id']}").get_data(as_text=True)
    assert "Window screen 35-1/2" in ov
    # keep on file: window screen line had no unit -> save it for unit 310
    r = post(c, f"/customers/{tid}/history/onfile", {"order_id": o["id"], "idx": 0, "unit": "310"})
    assert r.status_code == 302
    j = c.get(f"/api/unit?customer={tid}&form=window_screen&unit=310").get_json()
    assert j["blocks"] and "bronze" in str(j["blocks"][0].get("frame_finish", "")).lower()
