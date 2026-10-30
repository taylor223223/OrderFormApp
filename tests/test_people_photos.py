"""v2.14: people + positions saved to property profiles from the CRM; photo reference pages on order PDFs.
All names and data are made up."""
import io
import json
import os
import re
import tempfile

import pytest

os.environ.setdefault("ORDERAPP_DATA", tempfile.mkdtemp())

from orderapp import create_app  # noqa: E402


@pytest.fixture(scope="module")
def app():
    d = tempfile.mkdtemp()
    a = create_app(db_path=os.path.join(d, "p.db"), testing=True)
    with a.app_context():
        from orderapp.db import set_setting, x
        set_setting("output_dir", os.path.join(d, "out"))
        a.cid = x("INSERT INTO customers(name, address, city, state, zip, status) VALUES "
                  "('Cactus Bloom Apartments', '1 Test Rd', 'Phoenix', 'AZ', '85001', 'Active')")
        a.cid2 = x("INSERT INTO customers(name, status) VALUES ('Mesa Ridge Homes', 'Active')")
    return a


@pytest.fixture(scope="module")
def c(app):
    cl = app.test_client()

    def tok(path):
        return re.search(r'name="csrf-token" content="([^"]+)"', cl.get(path).get_data(as_text=True)).group(1)
    cl.post("/setup", data={"username": "tester", "password": "pw-123456", "password2": "pw-123456",
                            "csrf_token": tok("/setup")})
    cl.t = tok("/")
    return cl


def _people(app, cid):
    with app.app_context():
        from orderapp.db import q
        return {r["name"]: r["role"] for r in q("SELECT name, role FROM contacts WHERE customer_id=?", (cid,))}


def test_clean_role_and_matching(app):
    from orderapp.people import clean_role
    assert clean_role("maint super") == "Maint. Super"
    assert clean_role("Maintenance Supervisor") == "Maint. Super"
    assert clean_role("PM") == "Property Manager"
    assert clean_role("leasing") == "Office"
    assert clean_role("Regional VP") == "Regional VP"
    with app.app_context():
        from orderapp.people import save_person
        a = save_person(app.cid2, "Gus", "maintenance")
        b = save_person(app.cid2, "gus", "Maint. Super")           # same person, position updated
        cc = save_person(app.cid2, "Gus Larsen", "")               # fuller name fills in, no duplicate
        assert a == b == cc
    assert _people(app, app.cid2) == {"Gus Larsen": "Maint. Super"}


def test_log_people_with_positions(app, c):
    page = c.get("/crm/contacts").get_data(as_text=True)
    assert 'id="peopleRows"' in page and "Maint. Super" in page and "Property Manager" in page
    r = c.post("/crm/contacts/add", data={
        "csrf_token": c.t, "customer_id": app.cid, "kind": "In person", "temp": "Warm",
        "person_name": ["Dana Fry", "Lou", ""], "person_role": ["Maint. Super", "Office", ""],
        "person_phone": ["602-555-0101", "", ""], "subject": "Door check"})
    assert r.status_code == 302
    assert _people(app, app.cid) == {"Dana Fry": "Maint. Super", "Lou": "Office"}
    with app.app_context():
        from orderapp.db import q
        acts = q("SELECT contact_name, contact_id FROM activities WHERE customer_id=? ORDER BY id", (app.cid,))
        assert [a["contact_name"] for a in acts] == ["Dana Fry", "Lou"] and all(a["contact_id"] for a in acts)
        assert q("SELECT phone FROM contacts WHERE name='Dana Fry'", one=True)["phone"] == "602-555-0101"
    # suggestions for that property
    j = c.get(f"/crm/people/{app.cid}").get_json()
    assert {p["name"]: p["role"] for p in j["people"]} == {"Dana Fry": "Maint. Super", "Lou": "Office"}
    # shows on the profile with the position
    prof = c.get(f"/customers/{app.cid}").get_data(as_text=True)
    assert "Dana Fry" in prof and "Maint. Super" in prof and "last talked" in prof
    # log page shows the position column
    log = c.get("/crm/contacts?from=2000-01-01&to=2100-01-01").get_data(as_text=True)
    assert "<th>Position</th>" in log and 'value="Maint. Super"' in log


def test_inline_edit_saves_person(app, c):
    with app.app_context():
        from orderapp.db import x
        aid = x("INSERT INTO activities(customer_id, kind, occurred_at, created) VALUES (?, 'Call', '2030-01-02 10:00', 'x')",
                (app.cid,))
    r = c.post(f"/crm/contacts/{aid}", data={"csrf_token": c.t, "contact_name": "Rita Vance", "role": "pm",
                                              "kind": "Call", "date": "2030-01-02"},
               headers={"X-Requested-With": "fetch"})
    assert r.get_json()["ok"]
    assert _people(app, app.cid)["Rita Vance"] == "Property Manager"


def test_quick_log_saves_position(app, c):
    r = c.post("/crm/log", data={"csrf_token": c.t, "customer_id": app.cid, "kind": "Call",
                                 "contact_name": "Lou", "contact_role": "Office", "save_contact": "1"})
    assert r.status_code == 302
    assert list(_people(app, app.cid)).count("Lou") == 1


def test_old_names_backfilled(app, c):
    with app.app_context():
        from orderapp.db import x
        x("INSERT INTO activities(customer_id, kind, occurred_at, contact_name, created) "
          "VALUES (?, 'Text', '2030-01-03 10:00', 'Theo', 'x')", (app.cid2,))
    c.get("/crm/contacts")
    assert "Theo" in _people(app, app.cid2)


def _jpeg(w, h, color=(200, 80, 40)):
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (w, h), color).save(b, "JPEG", quality=95)
    return b.getvalue()


def test_photo_kept_full_quality():
    from orderapp.photos import shrink
    raw = _jpeg(1600, 1200)
    out, ext = shrink(raw)
    assert out == raw and ext == ".jpg"                # already fine -> stored exactly as taken
    big, _ = shrink(_jpeg(5000, 3000))
    from PIL import Image
    assert max(Image.open(io.BytesIO(big)).size) == 3200


def test_photos_on_order_pdf(app, c):
    payload = {"header": {"name": "Cactus Bloom Apartments", "po": "PO-77", "date": "10/15/2030"},
               "blocks": [{"qty": "1", "unit": "101"}], "comments": ""}
    r = c.post("/orders/save", data={"csrf_token": c.t, "form_key": "door", "customer_id": app.cid,
                                     "payload": json.dumps(payload), "action": "pdf",
                                     "photo_caption": "Unit 101 hall closet",
                                     "photos": [(io.BytesIO(_jpeg(1600, 1200)), "wide.jpg"),
                                                (io.BytesIO(_jpeg(900, 1600, (40, 90, 200))), "tall.jpg")]},
               content_type="multipart/form-data")
    assert r.status_code == 302
    oid = int(r.headers["Location"].rsplit("/", 1)[1])
    with app.app_context():
        from orderapp.db import q
        o = q("SELECT pdf_path FROM orders WHERE id=?", (oid,), one=True)
        assert q("SELECT COUNT(*) n FROM order_photos WHERE order_id=?", (oid,), one=True)["n"] == 2
    import pymupdf as fitz
    from orderapp.catalog import FORMS  # noqa: F401
    doc = fitz.open(o["pdf_path"])
    texts = [p.get_text() for p in doc]
    photo_pages = [i for i, t in enumerate(texts) if "PHOTO REFERENCE" in t]
    assert len(photo_pages) == 2 and photo_pages[-1] == doc.page_count - 1
    first = doc[photo_pages[0]]
    assert first.rect.width > first.rect.height                    # wide photo -> sideways page
    assert doc[photo_pages[1]].rect.width < doc[photo_pages[1]].rect.height
    assert "Unit 101 hall closet" in texts[photo_pages[0]] and "Photo 1 of 2" in texts[photo_pages[0]]
    img = first.get_images(full=True)[0]
    assert (img[2], img[3]) == (1600, 1200)                         # embedded at full size
    doc.close()
    # removing a photo re-makes the PDF
    with app.app_context():
        from orderapp.db import q
        pid = q("SELECT id FROM order_photos WHERE order_id=? ORDER BY id LIMIT 1", (oid,), one=True)["id"]
    c.post(f"/photos/{pid}/delete", data={"csrf_token": c.t})
    doc = fitz.open(o["pdf_path"])
    assert sum("PHOTO REFERENCE" in p.get_text() for p in doc) == 1
    doc.close()
    # editor shows the photo card + existing photos; send page says they're in the PDF
    ed = c.get(f"/orders/editor?order={oid}").get_data(as_text=True)
    assert "Photo reference" in ed and 'enctype="multipart/form-data"' in ed and "/photos/" in ed
    snd = c.get(f"/orders/{oid}/send").get_data(as_text=True)
    assert 'value="inpdf" checked' in snd


def test_previous_orders_with_us_style_dates(app):
    """An app order dated 10/15/2030 used to crash the editor's 'Previous orders' panel."""
    with app.app_context():
        from orderapp.db import x
        from orderapp.order_insights import previous_for
        x("INSERT INTO orders(customer_id, form_key, title, po, order_date, status, data, created) VALUES "
          "(?, 'door', 'Door Order Form', 'P9', '10/15/2030', 'Sent', ?, '2030-10-15 09:00')",
          (app.cid2, json.dumps({"blocks": [{"qty": "1", "width": "30", "height": "80"}]})))
        rows = previous_for(app.cid2, "door", [])
        assert rows and rows[0]["date"] == "2030-10-15"


def test_roller_shade_fabric_fields(app):
    from orderapp.catalog import FORMS, SHADE_FABRICS
    keys = [f["key"] for f in FORMS["roller_shade"]["blocks"][0]]
    assert {"shade_brand", "shade_fabric", "openness"} <= set(keys)
    assert "VX Solar Screen 3000" in SHADE_FABRICS and "Mermet E-Screen" in SHADE_FABRICS
    with app.app_context():
        from orderapp.pdf_fill import fill_form
        pdf, _ = fill_form("roller_shade", {"header": {"name": "Test Place"}, "blocks": [
            {"qty": "1", "unit": "5", "width": "36", "length": "60", "shade_brand": "Vertilux (VX)",
             "shade_fabric": "VX Solar Screen 3000", "openness": "3%", "color": "Other", "color_other": "White Pearl"}]})
    import pymupdf as fitz
    doc = fitz.open(stream=pdf, filetype="pdf")
    vals = {w.field_name: w.field_value for p in doc for w in p.widgets() or []}
    assert "VX Solar Screen 3000" in vals.values() and "3%" in vals.values() and "Vertilux (VX)" in vals.values()


def test_order_sizes_go_to_units_automatically(app, c):
    def save(payload, oid=""):
        r = c.post("/orders/save", data={"csrf_token": c.t, "form_key": "vertical_blind", "customer_id": app.cid2,
                                         "order_id": oid, "payload": json.dumps(payload), "action": "save"})
        return int(r.headers["Location"].rsplit("/", 1)[1])

    def meas():
        with app.app_context():
            from orderapp.db import q, loads
            return [(r["unit_number"], r["room"], loads(r["data"]).get("width"), loads(r["data"]).get("height"))
                    for r in q("""SELECT m.*, u.unit_number FROM measurements m JOIN units u ON u.id=m.unit_id
                                  WHERE m.customer_id=? ORDER BY u.unit_number, m.room, m.id""", (app.cid2,))]
    p = {"header": {"name": "Mesa Ridge Homes", "po": "A1"}, "blocks": [
        {"qty": "1", "unit": "201", "room": "Living Room", "width": "72", "height": "84"},
        {"qty": "1", "unit": "201", "room": "Bedroom", "width": "36", "height": "60"},
        {"qty": "1", "unit": "305", "room": "Living Room", "width": "96", "height": "84"}]}
    oid = save(p)
    assert sorted(meas()) == sorted([("201", "Bedroom", "36", "60"), ("201", "Living Room", "72", "84"),
                                     ("305", "Living Room", "96", "84")])
    save(dict(p, header={"name": "Mesa Ridge Homes", "po": "A2"}))      # same sizes on a 2nd order: no duplicates
    assert len(meas()) == 3
    p["blocks"][1]["width"] = "35-1/2"                                  # fix a typo on the first order
    save(p, oid)
    assert ("201", "Bedroom", "35-1/2", "60") in meas() and len(meas()) == 4   # 2nd order still has 36
    with app.app_context():
        from orderapp.suggest import measurements_for_unit
        _, _, rows = measurements_for_unit(app.cid2, "305", form_key="vertical_blind")
        assert rows
    assert "Sizes saved to units automatically" in c.get(f"/orders/{oid}").get_data(as_text=True)


def test_estimate_toggle_and_status_taps(app, c):
    payload = {"header": {"name": "Mesa Ridge Homes", "po": "Q-1", "date": "10/15/2030"},
               "blocks": [{"qty": "1", "unit": "9", "width": "30", "height": "80"}]}
    ed = c.get(f"/orders/editor?form=door&customer={app.cid2}").get_data(as_text=True)
    assert 'data-kind="quote"' in ed and 'name="kind"' in ed
    r = c.post("/orders/save", data={"csrf_token": c.t, "form_key": "door", "customer_id": app.cid2, "kind": "quote",
                                     "payload": json.dumps(payload), "action": "pdf"})
    oid = int(r.headers["Location"].rsplit("/", 1)[1])
    with app.app_context():
        from orderapp.db import q
        o = q("SELECT * FROM orders WHERE id=?", (oid,), one=True)
    assert o["kind"] == "quote" and o["title"] == "Door Quote Request" and o["status"] == "Ready"
    assert "Door Quote Request" in os.path.basename(o["pdf_path"])
    import pymupdf as fitz
    t = fitz.open(o["pdf_path"])[0].get_text()
    assert "Door Quote Form" in t and "Door Order Form" not in t        # the title bar says QUOTE
    page = c.get(f"/orders/{oid}").get_data(as_text=True)
    assert 'value="Quote ready" class="on"' in page and "<select name=\"status\">" not in page
    assert "Turn into an order" in page
    assert "Please price the attached" in c.get(f"/orders/{oid}/send").get_data(as_text=True)
    # one tap -> order sent: becomes an order, banner gone
    c.post(f"/orders/{oid}/status", data={"csrf_token": c.t, "tap": "Order sent"})
    with app.app_context():
        from orderapp.db import q
        o = q("SELECT * FROM orders WHERE id=?", (oid,), one=True)
    assert (o["kind"], o["status"], o["title"]) == ("order", "Sent", "Door Order Form") and o["sent_at"]
    assert "QUOTE REQUEST" not in fitz.open(o["pdf_path"])[0].get_text()
    c.post(f"/orders/{oid}/status", data={"csrf_token": c.t, "tap": "Delivered"})
    trk = c.get("/tracking?status=all").get_data(as_text=True)
    assert 'value="Delivered" class="on"' in trk and 'value="Order sent"' in trk


def test_holds_csv_file_upload(app, c):
    """Picking the hold CSV (no pasted text) used to read nothing - the header row wasn't first."""
    csv_text = ("Customer,Balance,Rep,Terms,Current,31-60,61-90,Over 90\n"
                "Mesa Ridge Homes,\"1,206.05\",SS,ACCOUNT HOLD,0.00,844.83,361.22,0.00\n")
    c.post("/holds/import", data={"csrf_token": c.t, "file": (io.BytesIO(("﻿" + csv_text).encode()), "holds.csv")},
           content_type="multipart/form-data")
    with app.app_context():
        from orderapp.db import q
        h = q("SELECT * FROM account_holds WHERE name='Mesa Ridge Homes'", one=True)
    assert h and h["balance"] == 1206.05 and h["customer_id"] == app.cid2


def test_cabinet_hinge_cuts(app):
    from orderapp.catalog import FORMS
    for fk in ("cabinet", "cabinet_custom", "cabinet_prefab"):
        keys = [f["key"] for f in FORMS[fk]["blocks"][0]]
        assert {"hinge_cut", "hinge_top", "hinge_bottom", "hinge_center", "hinge_cut_dia"} <= set(keys), fk
        assert len(FORMS[fk]["blocks"]) == 2                    # still two cabinets per page
    with app.app_context():
        from orderapp.pdf_fill import fill_form
        pdf, _ = fill_form("cabinet_custom", {"header": {"name": "Test Place"}, "blocks": [
            {"qty": "1", "unit": "7", "cab_kind": "Wall", "width": "30", "height": "30", "depth": "12",
             "hinge_cut": "Yes", "hinge_top": "3-1/2", "hinge_bottom": "3", "hinge_center": "15-1/2", "hinge_cut_dia": "35mm"}]})
    import pymupdf as fitz
    doc = fitz.open(stream=pdf, filetype="pdf")
    vals = {w.field_name: w.field_value for p in doc for w in p.widgets() or []}
    assert {"3-1/2", "3", "15-1/2"} <= set(vals.values())
    on = [n for n, v in vals.items() if v in ("Yes", True, "On") and "hinge_cut" in n.lower()]
    assert "L0.hinge_cut.yes" in on and "L0.hinge_cut_dia.35mm" in on, on


def test_door_edge_full_book_vs_beveled(app, c):
    from orderapp.door_edge import BEVELED, FULL_BOOK, guess, standard_size
    assert guess("30", "80") == FULL_BOOK and guess("36", "80") == FULL_BOOK
    assert guess("29-3/4", "79-1/2") == BEVELED and guess("35 3/4", "79 1/2") == BEVELED
    assert guess("31", "78") == "" and standard_size("30", "80") == "29-3/4 x 79-1/2"
    from orderapp.catalog import FORMS
    for fk in ("door", "new_door"):
        keys = [f["key"] for f in FORMS[fk]["blocks"][0]]
        assert keys.index("door_edge") == keys.index("height") + 1, fk
    assert "door_edge" not in [f["key"] for f in FORMS["prehung"]["blocks"][0]]   # pre-hung: no edge (10/26)
    # saving fills the edge from the size when it wasn't picked; a picked one is kept
    payload = {"header": {"name": "Mesa Ridge Homes", "po": "E-1"}, "blocks": [
        {"qty": "1", "unit": "11", "width": "30", "height": "80"},
        {"qty": "1", "unit": "12", "width": "35-3/4", "height": "79-1/2"},
        {"qty": "1", "unit": "13", "width": "30", "height": "80", "door_edge": BEVELED}]}
    r = c.post("/orders/save", data={"csrf_token": c.t, "form_key": "door", "customer_id": app.cid2,
                                     "payload": json.dumps(payload), "action": "pdf"})
    oid = int(r.headers["Location"].rsplit("/", 1)[1])
    with app.app_context():
        from orderapp.db import loads, q
        o = q("SELECT * FROM orders WHERE id=?", (oid,), one=True)
        blocks = loads(o["data"])["blocks"]
    assert [b["door_edge"] for b in blocks] == [FULL_BOOK, BEVELED, BEVELED]
    import pymupdf as fitz
    txt = " ".join(str(w.field_value) for p in fitz.open(o["pdf_path"]) for w in (p.widgets() or []))
    txt += " ".join(p.get_text() for p in fitz.open(o["pdf_path"]))
    assert "Full Book" in txt and "Beveled" in txt
    ed = c.get(f"/orders/editor?form=door&customer={app.cid2}").get_data(as_text=True)
    assert "function doorEdge" in ed and "Full Book size" in ed


def test_every_door_style_has_a_picture():
    from orderapp import catalog, option_photos
    for key, f in catalog.FORMS.items():
        for blk in f["blocks"][:1]:
            for fd in blk:
                if fd.get("photos") != "door_style":
                    continue
                for opt in fd["options"]:
                    label = opt[0] if isinstance(opt, (tuple, list)) else opt
                    if label in ("Other", "Avalon"):    # Avalon: no photo yet (Taylor can send one)
                        continue
                    assert option_photos.find("door_style", label), (key, label)
    for pn in catalog.DOOR_PANELS:
        assert option_photos.find("door_style", pn) or pn == "Flush", pn
    # panels + design together
    assert catalog.door_photo({"panels": "4 Panel", "style": "Shaker"}).endswith("4-panel-shaker.jpg")
    assert catalog.door_photo({"panels": "6 Panel", "style": "Colonial", "door_top": "Arched Top"}).endswith(
        "6-panel-colonial.jpg")
    assert catalog.door_photo({"panels": "2 Panel", "style": "", "door_top": "Arched Top"}).endswith(
        "2-panel-arch.jpg")


def test_export_page_and_files(c):
    from openpyxl import load_workbook
    r = c.get("/export")
    assert r.status_code == 200 and b"Customer list" in r.data
    r = c.get("/export/customers.csv?missing=address")
    txt = r.get_data(as_text=True)
    assert "Mesa Ridge Homes" in txt and "Cactus Bloom" not in txt
    assert txt.lstrip("﻿").startswith("Property,Account Type,Services,Acct #,Address,City,State,Zip,Mgmt Co")
    for name in ("customers", "contacts", "sizes", "orders", "sales", "activities", "tasks", "deals", "items", "holds"):
        r = c.get(f"/export/{name}.xlsx")
        assert r.status_code == 200, name
        load_workbook(io.BytesIO(r.data))
    wb = load_workbook(io.BytesIO(c.get("/export/everything.xlsx").data))
    assert len(wb.sheetnames) == 10
    assert c.get("/export/nope.csv").status_code == 404


def test_account_types_links_and_mgmt(app, c):
    from orderapp import accounts as A
    assert A.guess_type("Kinella Construction") == A.CONTRACTOR
    assert A.guess_type("Desert Bloom (Renovations)") == A.PROPERTY
    assert A.guess_type("Sundial Property Management") == A.MGMT
    assert A.guess_type("Painted Trails") == A.PROPERTY
    assert A.clean_type("Apartment Community") == A.PROPERTY and A.clean_type("vendor") == A.VENDOR
    # add a contractor and a mgmt company from the form, linked from a property
    r = c.post("/customers/new", data={"csrf_token": c.t, "name": "Sonoran Turn Crew LLC", "acct_type": "Contractor",
                                       "services": "Make-ready / turns", "link_to": str(app.cid)})
    assert r.status_code == 302
    r = c.post("/customers/new", data={"csrf_token": c.t, "name": "Saguaro Living Group", "acct_type": "Management Company"})
    with app.app_context():
        from orderapp.db import q, x
        crew = q("SELECT * FROM customers WHERE name='Sonoran Turn Crew LLC'", one=True)
        assert crew["acct_type"] == "Contractor" and crew["services"] == "Make-ready / turns"
        assert [l["name"] for l in A.links_for(app.cid)] == ["Sonoran Turn Crew LLC"]
        x("UPDATE customers SET mgmt='Saguaro Living Group' WHERE id=?", (app.cid,))
        mg = q("SELECT * FROM customers WHERE name='Saguaro Living Group'", one=True)
        assert [m["name"] for m in A.managed_by(mg)] == ["Cactus Bloom Apartments"]
    page = c.get(f"/customers/{crew['id']}").get_data(as_text=True)
    assert "Works with" in page and "Cactus Bloom Apartments" in page and "Floorplans" not in page
    assert "Properties they manage" in c.get(f"/customers/{mg['id']}").get_data(as_text=True)
    lst = c.get("/customers?type=Contractor").get_data(as_text=True)
    assert "Sonoran Turn Crew" in lst and "Mesa Ridge Homes" not in lst
    assert "+ Contractor" in c.get("/crm").get_data(as_text=True)
    # link by name, then remove
    r = c.post(f"/customers/{app.cid2}/links", data={"csrf_token": c.t, "other_name": "Sonoran Turn Crew LLC",
                                                     "note": "does their turns"})
    with app.app_context():
        assert len(A.links_for(crew["id"])) == 2
        lid = A.links_for(app.cid2)[0]["id"]
    c.post(f"/customers/{app.cid2}/links", data={"csrf_token": c.t, "remove": str(lid)})
    with app.app_context():
        assert len(A.links_for(crew["id"])) == 1
    txt = c.get("/export/customers.csv?type=Contractor").get_data(as_text=True)
    assert "Sonoran Turn Crew LLC,Contractor,Make-ready / turns" in txt and "Cactus" not in txt


def test_reorder_from_purchase_history(app, c):
    import fitz
    with app.app_context():
        from orderapp.db import q, x
        sid = x("INSERT INTO sales(invoice_no, invoice_date, customer_id, customer_name, total) VALUES "
                "('90001', '2026-08-10', ?, 'Cactus Bloom Apartments', 300)", (app.cid,))
        x("INSERT INTO sales_lines(sale_id, category, item_no, description, qty, unit_price, amount) VALUES "
          "(?, 'Horizontal Blinds', 'HB-2', '2\" Faux wood blind 35 x 48', 4, 30, 120)", (sid,))
        x("INSERT INTO sales_lines(sale_id, category, item_no, description, qty, unit_price, amount) VALUES "
          "(?, 'Doors', 'D-30', '30x80 HC Colonist LH', 2, 90, 180)", (sid,))
    page = c.get(f"/customers/{app.cid}?tab=products").get_data(as_text=True)
    assert "Make reorder form" in page and 'name="sel"' in page and "90001" in page
    lines_view = c.get(f"/customers/{app.cid}?tab=products&view=lines").get_data(as_text=True)
    assert 'name="inv_0" value="90001"' in lines_view and "Tick every line on this invoice" in lines_view
    # nothing ticked -> back with a message
    r = c.post(f"/customers/{app.cid}/reorder", data={"csrf_token": c.t})
    assert r.status_code == 302 and "tab=products" in r.headers["Location"]
    r = c.post(f"/customers/{app.cid}/reorder", data={
        "csrf_token": c.t, "sel": ["0", "1"], "qty_0": "4", "desc_0": '2" Faux wood blind 35 x 48', "item_0": "HB-2", "inv_0": "90001",
        "size_0": "35 x 48", "last_0": "2026-08-10", "price_0": "30.00", "qty_1": "1", "desc_1": "30x80 HC Colonist LH",
        "item_1": "D-30", "po": "PO-77", "contact": "Pat - Maint. Super", "comments": "Deliver to the shop"})
    assert r.status_code == 302
    oid = int(r.headers["Location"].rstrip("/").split("/")[-1])
    with app.app_context():
        o = q("SELECT * FROM orders WHERE id=?", (oid,), one=True)
        assert o["form_key"] == "reorder" and o["status"] == "Ready" and o["po"] == "PO-77"
        txt = "".join(p.get_text() for p in fitz.open(o["pdf_path"]))
    assert "#90001" in txt and "REORDER REQUEST" in txt and "HB-2" in txt and "Faux wood" in txt and "PO-77" in txt and "Pat - Maint. Super" in txt
    view = c.get(f"/orders/{oid}").get_data(as_text=True)
    assert "Reorder lines" in view and "Email it" in view
    assert c.get(f"/orders/{oid}/pdf").status_code == 200
    assert c.get(f"/orders/{oid}/send").status_code == 200
    # edit: drop the door, add a line
    r = c.post(f"/orders/{oid}/reorder_lines", data={
        "csrf_token": c.t, "n": "2", "po": "PO-78", "qty_0": "6", "desc_0": '2" Faux wood blind 35 x 48', "item_no_0": "HB-2",
        "qty_1": "", "desc_1": "30x80 HC Colonist LH", "qty_2": "3", "desc_2": "Door stop", "item_no_2": "DS-1"})
    assert r.status_code == 302
    with app.app_context():
        from orderapp.db import loads
        d = loads(q("SELECT data FROM orders WHERE id=?", (oid,), one=True)["data"])
    assert [(l["qty"], l["item_no"]) for l in d["lines"]] == [("6", "HB-2"), ("3", "DS-1")] and d["header"]["po"] == "PO-78"
    assert "Reorder Request" in c.get("/tracking").get_data(as_text=True)


def test_plan_week_by_area_and_accounts(app, c):
    from datetime import date, timedelta
    with app.app_context():
        from orderapp.db import x, q
        old = (date.today() - timedelta(days=70)).isoformat() + " 09:00"
        mk = lambda n, city, zp, lat, lon: x("INSERT INTO customers(name, address, city, state, zip, status, lat, lon) "
                                             "VALUES (?,?,?,?,?,'Active',?,?)", (n, '1 Test St', city, 'AZ', zp, lat, lon))
        south = mk("Ocotillo South Flats", "Phoenix", "85040", 33.395, -112.070)
        south2 = mk("Baseline Palms", "Phoenix", "85041", 33.380, -112.080)
        scott = mk("Camelback Vista Scottsdale", "Scottsdale", "85251", 33.495, -111.925)
        for cid in (south, south2, scott):
            x("INSERT INTO activities(customer_id, kind, occurred_at, subject) VALUES (?,?,?,?)", (cid, "Call", old, "x"))
        sid = x("INSERT INTO sales(invoice_no, invoice_date, customer_id, customer_name, total) VALUES "
                "('91001', ?, ?, 'Baseline Palms', 2500)", ((date.today() - timedelta(days=40)).isoformat(), south2))
    mon = date.today() - timedelta(days=date.today().weekday())
    wk = [mon + timedelta(days=i) for i in range(5)]
    day = next((d for d in wk if d >= date.today()), wk[-1]).isoformat()
    from orderapp import plan_focus as PF
    with app.app_context():
        assert PF.resolve_area("South Central Phoenix")[3] == "area"
        assert PF.resolve_area("85251")[3] == "zip" and PF.resolve_area("scottsdale")[3] == "city"
    r = c.post("/plan/area", data={"csrf_token": c.t, "day": day, "area": "South Central Phoenix"})
    assert r.status_code == 302
    page = c.get("/plan").get_data(as_text=True)
    assert "around South Central Phoenix" in page
    with app.app_context():
        from orderapp import planner as PL
        recs = PL.recommendations(7)
        work = [d.isoformat() for d in wk if d.isoformat() >= date.today().isoformat()] or [d.isoformat() for d in wk]
        sugg = PF.place_recs(recs, work)
        by = {r["name"]: r for r in recs}
        assert by["Ocotillo South Flats"]["day"] == day and "mi from South Central Phoenix" in by["Ocotillo South Flats"]["near"]
        assert by["Camelback Vista Scottsdale"]["day"] != day or by["Camelback Vista Scottsdale"].get("far")
    # add an account to the day, then look around it
    r = c.post("/plan/day_add", data={"csrf_token": c.t, "day": day, "account": "Ocotillo South Flats"})
    assert r.status_code == 302
    with app.app_context():
        from orderapp.db import q
        assert q("SELECT 1 FROM route_stops WHERE day=? AND customer_id=?", (day, south), one=True)
    near = c.get("/plan/near?account=Ocotillo South Flats&radius=3").get_data(as_text=True).split("<table")[-1]
    assert "Baseline Palms" in near and "Camelback Vista Scottsdale" not in near
    near = c.get("/plan/near?area=Scottsdale&radius=3").get_data(as_text=True).split("<table")[-1]
    assert "Camelback Vista Scottsdale" in near and "Baseline Palms" not in near
    r = c.post("/plan/add", data={"csrf_token": c.t, "pick": "1", "cid_1": str(south2), "name_1": "Baseline Palms",
                                  "why_1": "Near", "day_1": day, "back": "/plan/near?account=x"})
    assert r.status_code == 302
    c.post("/plan/area", data={"csrf_token": c.t, "day": day, "area": ""})


def test_territory_filter(app, c):
    from datetime import date, timedelta
    from orderapp import territory as T
    with app.app_context():
        from orderapp.db import x, set_setting
        set_setting("report_rep", "Taylor Anderson")
        old = (date.today() - timedelta(days=80)).isoformat() + " 09:00"
        inside = x("INSERT INTO customers(name, zip, city, status, lat, lon) VALUES ('Inside Terr Flats','85013','Phoenix','Active',33.51,-112.08)")
        outside = x("INSERT INTO customers(name, zip, city, status, lat, lon) VALUES ('Mesa Far Away','85210','Mesa','Active',33.39,-111.84)")
        tempe_s = x("INSERT INTO customers(name, zip, city, status, lat, lon) VALUES ('Tempe South 202','85281','Tempe','Active',33.41,-111.93)")
        tempe_n = x("INSERT INTO customers(name, zip, city, status, lat, lon) VALUES ('Tempe North 202','85281','Tempe','Active',33.45,-111.93)")
        for cid in (inside, outside, tempe_s, tempe_n):
            x("INSERT INTO activities(customer_id, kind, occurred_at, subject) VALUES (?,?,?,?)", (cid, "Call", old, "x"))
        ids = T.my_ids()
        assert inside in ids and tempe_n in ids and outside not in ids and tempe_s not in ids
        from orderapp import planner as PL
        names = [r["name"] for r in PL.recommendations(7)]
        assert "Inside Terr Flats" in names and "Mesa Far Away" not in names and "Tempe South 202" not in names
    # Settings page saves ZIPs; Mesa added to my list -> it shows up
    page = c.get("/settings/territory").get_data(as_text=True)
    assert "Taylor Anderson" in page and "85013" in page
    r = c.post("/settings/territory", data={"csrf_token": c.t, "zips_0": T.DEFAULT_MY_ZIPS + " 85210", "only": "1",
                                            "include_billed": "1", "partial": "85281 85288"})
    assert r.status_code == 302
    with app.app_context():
        assert outside in T.my_ids()
    lst = c.get("/customers?terr=mine").get_data(as_text=True)
    assert "Inside Terr Flats" in lst and "Tempe South 202" not in lst
    c.post("/settings/territory", data={"csrf_token": c.t, "zips_0": T.DEFAULT_MY_ZIPS, "only": "1",
                                        "include_billed": "1", "partial": "85281 85288"})
